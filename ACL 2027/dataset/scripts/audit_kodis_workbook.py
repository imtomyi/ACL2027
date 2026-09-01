#!/usr/bin/env python3
"""Create aggregate, non-text KODIS receipt metrics from the received workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


LINE_RE = re.compile(r"^(?P<time>nan|\d+)\s+(?P<role>Buyer|Seller):\s*(?P<text>.*)$")
WORD_RE = re.compile(r"\b[\w']+\b", re.UNICODE)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def percentile(values: list[int], proportion: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * proportion
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return float(ordered[lower])
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def distribution(values: list[int]) -> dict[str, float | int | None]:
    return {
        "n": len(values),
        "sum": sum(values),
        "min": min(values) if values else None,
        "p25": percentile(values, 0.25),
        "median": statistics.median(values) if values else None,
        "p75": percentile(values, 0.75),
        "max": max(values) if values else None,
        "mean": statistics.fmean(values) if values else None,
    }


def is_true(value: object) -> bool:
    return value is True or (isinstance(value, str) and value.strip().lower() == "true")


def parse_messages(chat: str) -> tuple[list[tuple[str, str]], int]:
    messages: list[tuple[str, str]] = []
    unparsable_nonblank_lines = 0
    for line in chat.splitlines():
        line = line.strip()
        if not line:
            continue
        match = LINE_RE.match(line)
        if match:
            messages.append((match.group("role").lower(), match.group("text").strip()))
        elif messages:
            role, prior = messages[-1]
            messages[-1] = (role, f"{prior} {line}".strip())
        else:
            unparsable_nonblank_lines += 1
    return messages, unparsable_nonblank_lines


def audit(workbook_path: Path) -> tuple[dict, list[dict]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    headers[0] = headers[0] or "dyad_id"
    index = {str(value): position + 1 for position, value in enumerate(headers)}
    required = {"dyad_id", "formattedChat", "Outcome", "buyer_is_AI", "seller_is_AI"}
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    lane_counts: Counter[str] = Counter()
    outcome_by_lane: dict[str, Counter[str]] = {}
    transcript_counts: dict[str, list[int]] = {}
    word_counts: dict[str, list[int]] = {}
    char_counts: dict[str, list[int]] = {}
    first_role: dict[str, Counter[str]] = {}
    structure: dict[str, Counter[str]] = {}
    countries: dict[str, Counter[str]] = {"buyer": Counter(), "seller": Counter()}
    same_country_human_human = 0
    different_country_human_human = 0
    candidate_dialogues: Counter[str] = Counter()
    dyad_ids: list[str] = []
    exact_chat_counts: Counter[str] = Counter()
    nulls: Counter[str] = Counter()

    for row in range(2, sheet.max_row + 1):
        values = {str(header): sheet.cell(row, column).value for column, header in enumerate(headers, 1)}
        dyad_id = values["dyad_id"]
        dyad_ids.append(str(dyad_id))
        buyer_ai = is_true(values["buyer_is_AI"])
        seller_ai = is_true(values["seller_is_AI"])
        lane = "human_human" if not buyer_ai and not seller_ai else "human_ai" if buyer_ai ^ seller_ai else "ai_ai"
        lane_counts[lane] += 1
        outcome = str(values["Outcome"]) if values["Outcome"] is not None else "missing"
        outcome_by_lane.setdefault(lane, Counter())[outcome] += 1
        for header, value in values.items():
            if value is None or (isinstance(value, str) and not value.strip()):
                nulls[header] += 1

        chat = values["formattedChat"]
        structure.setdefault(lane, Counter())
        first_role.setdefault(lane, Counter())
        transcript_counts.setdefault(lane, [])
        word_counts.setdefault(lane, [])
        char_counts.setdefault(lane, [])
        if not isinstance(chat, str) or not chat.strip():
            structure[lane]["missing_transcript"] += 1
            continue
        exact_chat_counts[hashlib.sha256(chat.encode("utf-8")).hexdigest()] += 1
        messages, unparsable = parse_messages(chat)
        transcript_counts[lane].append(len(messages))
        words = sum(len(WORD_RE.findall(text)) for _, text in messages)
        word_counts[lane].append(words)
        char_counts[lane].append(sum(len(text) for _, text in messages))
        structure[lane]["unparsable_leading_line_dialogues"] += int(unparsable > 0)
        structure[lane]["timestamp_nan_dialogues"] += int(any(line.startswith("nan ") for line in chat.splitlines()))
        if messages:
            roles = [role for role, _ in messages]
            first_role[lane][roles[0]] += 1
            structure[lane]["under_8_messages"] += int(len(messages) < 8)
            structure[lane]["nonalternating_roles"] += int(any(roles[i] == roles[i - 1] for i in range(1, len(roles))))
            structure[lane]["missing_buyer_or_seller"] += int(set(roles) != {"buyer", "seller"})
            if len(messages) >= 8 and roles[0] == "buyer" and set(roles) == {"buyer", "seller"} and not any(roles[i] == roles[i - 1] for i in range(1, len(roles))):
                candidate_dialogues[lane] += 1
        else:
            structure[lane]["zero_parsed_messages"] += 1

        if lane == "human_human":
            for role, field in (("buyer", "b_country"), ("seller", "s_country")):
                value = values.get(field)
                countries[role][str(value).strip() if value is not None else "missing"] += 1
            buyer_country = values.get("b_country")
            seller_country = values.get("s_country")
            if buyer_country is not None and seller_country is not None:
                if str(buyer_country).strip() == str(seller_country).strip():
                    same_country_human_human += 1
                else:
                    different_country_human_human += 1

    duplicate_groups = [count for count in exact_chat_counts.values() if count > 1]
    metrics = {
        "audit_version": "1.0",
        "source_file": workbook_path.name,
        "source_sha256": sha256(workbook_path),
        "worksheet": sheet.title,
        "rows": sheet.max_row - 1,
        "columns": sheet.max_column,
        "unique_dyad_ids": len(set(dyad_ids)),
        "duplicate_dyad_id_rows": len(dyad_ids) - len(set(dyad_ids)),
        "interaction_counts": dict(lane_counts),
        "outcome_counts_by_interaction": {key: dict(value) for key, value in outcome_by_lane.items()},
        "transcript_message_count": {key: distribution(value) for key, value in transcript_counts.items()},
        "transcript_word_count": {key: distribution(value) for key, value in word_counts.items()},
        "transcript_character_count": {key: distribution(value) for key, value in char_counts.items()},
        "first_parsed_role": {key: dict(value) for key, value in first_role.items()},
        "transcript_structure_flags": {key: dict(value) for key, value in structure.items()},
        "dialogues_passing_existing_structural_rules": dict(candidate_dialogues),
        "exact_duplicate_transcript_groups": len(duplicate_groups),
        "rows_in_exact_duplicate_transcript_groups": sum(duplicate_groups),
        "column_missing_counts": dict(sorted(nulls.items())),
        "schema_findings": {
            "one_row_per_dyad": True,
            "transcript_storage": "formattedChat multiline string with timestamp-or-nan, role, and text",
            "explicit_interaction_marker": "buyer_is_AI and seller_is_AI",
            "explicit_completion_marker_present": False,
            "explicit_event_type_present": False,
            "stable_cross_dialogue_participant_id_present": False,
            "message_id_present": False,
            "reply_to_present": False,
        },
        "human_human_country_summary": {
            "buyer_distinct_values_including_missing": len(countries["buyer"]),
            "seller_distinct_values_including_missing": len(countries["seller"]),
            "same_country_dyads_with_both_observed": same_country_human_human,
            "different_country_dyads_with_both_observed": different_country_human_human,
        },
    }
    country_rows = []
    for role, counts in countries.items():
        total = sum(counts.values())
        for country, count in counts.most_common():
            country_rows.append({"role": role, "country": country, "dyads": count, "share": count / total if total else None})
    return metrics, country_rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    metrics, countries = audit(args.workbook)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "aggregate_metrics.json").write_text(json.dumps(metrics, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    with (args.output_dir / "human_human_country_distribution.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=["role", "country", "dyads", "share"])
        writer.writeheader()
        writer.writerows(countries)


if __name__ == "__main__":
    main()
